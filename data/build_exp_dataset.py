# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""
Build the pair-level experimental dataset: one row per photocatalyst-quencher pair.

Reads   experimental_data_processed/stern_volmer_fits.csv
        experimental_data_processed/photocatalyst_properties_exp.xlsx
        experimental_data_processed/quencher_properties_exp.xlsx
        compound_identification.xlsx   (ids and names)
Writes  exp_dataset.csv     input to build_exp_dataset_ent.py
        exp_dataset.xlsx    the same table in the property-workbook layout

    uv run data/build_exp_dataset.py

Labels
    `quenches` is taken from stern_volmer_fits.csv. It is set on K_SV, since a cutoff on
    k_q would shift with tau0. k_q = K_SV / tau0.
    A pair that does not quench is left-censored: `censored = 1`, and its fitted K_SV is
    kept.

Only measured values are included. The energy-transfer subset, with calculated
properties and driving forces, is built by build_exp_dataset_ent.py.

Photocatalysts
    `pc_new_id` (1-27), `pc_name` is PC<new_id>. All 27 are included. PC22 is excluded by
    the theory benchmark (see SI), but kept here.
"""

import csv
import sys
from pathlib import Path

DATA = Path(__file__).resolve().parent
sys.path.insert(0, str(DATA))

from openpyxl import load_workbook  # noqa: E402

from compound_identification import PHOTOCATALYSTS  # noqa: E402
from properties_workbook import Block, Column, write_workbook  # noqa: E402

EXPERIMENTAL = DATA / "experimental_data_processed"
OUT_CSV = DATA / "exp_dataset.csv"
OUT_XLSX = DATA / "exp_dataset.xlsx"

# Row 1 of a property workbook names the blocks; headers are on row 2.
WORKBOOK_HEADER_ROW = 2

SIGFIGS = 6

# (block, [(column, description, source)]) in table order.
LAYOUT = [
    ("Identity", [
        ("pc_new_id", "Photocatalyst id, 1-27.", "photocatalyst_properties_exp.xlsx: new_id"),
        ("pc_name", "PC<pc_new_id>.", "computed here"),
        ("quencher_id", "Quencher number.", "quencher_properties_exp.xlsx"),
        ("quencher_name", "Q<n>.", "quencher_properties_exp.xlsx"),
    ]),
    ("Measurement", [
        ("K_SV / M-1", "Stern-Volmer constant (fitted slope).", "stern_volmer_fits.csv: K_SV"),
        ("tau0 / ns", "Unquenched lifetime of the photocatalyst.",
         "photocatalyst_properties_exp.xlsx"),
        ("k_q / M-1 s-1", "Quenching rate constant, K_SV / tau0.", "computed here"),
    ]),
    ("Ground truth", [
        ("quenches", "1 if the fit meets one of the (K_SV, r2) thresholds.",
         "stern_volmer_fits.csv: quenches"),
        ("censored", "1 where quenches = 0: the true K_SV is below what was detectable.",
         "computed here"),
    ]),
]
COLUMNS = [name for _, columns in LAYOUT for name, _, _ in columns]


def sig(x, n=SIGFIGS):
    if x in (None, ""):
        return ""
    return 0.0 if float(x) == 0 else float(f"{float(x):.{n}g}")


def load(path, key):
    """A property workbook as {key: row}, blanks as "". Exits if the file is missing."""
    if not path.exists():
        sys.exit(f"missing input: {path}\nrun the data stage first: uv run reproduce.py data")
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(min_row=WORKBOOK_HEADER_ROW, values_only=True)
    header = next(rows)
    table = {}
    for values in rows:
        row = {h: ("" if v is None else v) for h, v in zip(header, values) if h}
        if row.get(key) not in ("", None):
            table[int(row[key])] = row
    return table


def number(row, field):
    """A numeric field, or None when absent or blank."""
    if row is None:
        return None
    value = row.get(field, "")
    return float(value) if value not in ("", None) else None


def main():
    fits_path = EXPERIMENTAL / "stern_volmer_fits.csv"
    if not fits_path.exists():
        sys.exit(f"missing input: {fits_path}\nrun the data stage first")

    pc_exp = load(EXPERIMENTAL / "photocatalyst_properties_exp.xlsx", "new_id")
    q_exp = load(EXPERIMENTAL / "quencher_properties_exp.xlsx", "quencher_id")

    rows = []
    with fits_path.open(encoding="utf-8", newline="") as fh:
        for fit in csv.DictReader(fh):
            pc_id, q_id = int(fit["pc_new_id"]), int(fit["quencher_id"])
            if pc_id not in PHOTOCATALYSTS:
                sys.exit(f"Stern-Volmer fits for PC{pc_id}, which is not in the data")
            pce, qe = pc_exp.get(pc_id), q_exp.get(q_id)

            tau0 = number(pce, "tau0 / ns")
            k_sv = number(fit, "K_SV")
            # tau0 is in ns; k_q comes out in M-1 s-1.
            k_q = k_sv / (tau0 * 1e-9) if (k_sv is not None and tau0) else None
            quenches = fit["quenches"] == "1"

            rows.append({
                "pc_new_id": pc_id,
                "pc_name": f"PC{pc_id}",
                "quencher_id": q_id,
                "quencher_name": (qe or {}).get("quencher_name", ""),
                "K_SV / M-1": sig(k_sv),
                "tau0 / ns": sig(tau0),
                "k_q / M-1 s-1": sig(k_q),
                "quenches": int(quenches),
                "censored": int(not quenches),
            })

    rows.sort(key=lambda row: (row["pc_new_id"], row["quencher_id"]))
    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    blocks = [Block(title, [Column(name, description, source,
                                   lambda i, name=name: rows[i][name] if rows[i][name] != "" else None)
                            for name, description, source in columns])
              for title, columns in LAYOUT]
    write_workbook(OUT_XLSX, "Pairs", list(range(len(rows))), blocks, """
        One row per photocatalyst-quencher pair, sorted by pc_new_id then quencher_id.
        Built by data/build_exp_dataset.py; exp_dataset.csv holds the same table.
        Measured values and labels only. exp_dataset_ent adds calculated properties.
        PC22 is included; the theory benchmark excludes it (see SI).
    """)

    n_pc = len({r["pc_new_id"] for r in rows})
    n_q = len({r["quencher_id"] for r in rows})
    print(f"{len(rows)} pairs ({n_pc} photocatalysts x {n_q} quenchers) -> "
          f"{OUT_CSV.name}, {OUT_XLSX.name}")
    print(f"  quenches {sum(r['quenches'] for r in rows)}, "
          f"censored {sum(r['censored'] for r in rows)}")
    for name in COLUMNS:
        missing = sum(1 for r in rows if r[name] == "")
        if missing:
            print(f"  {name}: blank for {missing} pairs")


if __name__ == "__main__":
    main()
