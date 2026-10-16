# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""
Molar extinction coefficients of the quenchers from a plate-reader absorption scan.

Reads   ../experimental_data_raw/quencher_absorption/
        2025-12-22 10-53-28 Absorption_Scan_360_650_56well_published.csv
        quenchers_deepwell.xlsx     plate map: quencher, well, stock concentration
Writes  ../experimental_data_processed/absorption_extinction.csv

    uv run data/experimental_data_scripts/absorption_extinction.py

epsilon = (A * 10 mm / 5.846 mm - mean A of the solvent wells, column 7) / c_stock.
Stern-Volmer constants are not inner-filter corrected.
"""

import csv
import re
import statistics
import sys
from pathlib import Path

from openpyxl import load_workbook

RAW = Path(__file__).resolve().parent.parent / "experimental_data_raw" / "quencher_absorption"
OUT = Path(__file__).resolve().parent.parent / "experimental_data_processed" / "absorption_extinction.csv"

SCAN = "2025-12-22 10-53-28 Absorption_Scan_360_650_56well_published.csv"
PLATE_MAP = "quenchers_deepwell.xlsx"

# Path length through the deep-well plate; readings are scaled to a 10 mm path.
PATH_LENGTH_MM = 5.846
REFERENCE_PATH_MM = 10.0

BLANK_WELLS = [f"{row}7" for row in "ABCDEFGH"]   # column 7 is solvent only

WAVELENGTHS = [380, 390, 400, 405, 415, 425, 450, 475,
               490, 495, 500, 510, 520, 580, 590, 600, 625, 645]

SIGFIGS = 6


def sig(x, n=SIGFIGS):
    if x is None:
        return ""
    return 0.0 if x == 0 else float(f"{float(x):.{n}g}")


def normalise_well(value):
    """'A07' -> 'A7'."""
    s = str(value).strip()
    m = re.fullmatch(r"([A-Za-z]+)\s*0*(\d+)", s)
    return f"{m.group(1).upper()}{int(m.group(2))}" if m else s.upper()


def read_scan(path):
    """well -> {wavelength_nm: mean absorbance}, and the relative spread (%) of repeat reads.

    Comma-delimited or semicolon-delimited (decimal commas) plate-reader exports.
    """
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except UnicodeDecodeError:
        lines = path.read_text(encoding="latin-1").splitlines()

    header_idx = next((i for i, ln in enumerate(lines) if "Wavelength [nm]" in ln), None)
    if header_idx is None:
        raise ValueError(f"no 'Wavelength [nm]' header in {path.name}")

    header_line = lines[header_idx]
    delimiter = ";" if header_line.count(";") > header_line.count(",") else ","
    header = [f.strip() for f in header_line.split(delimiter)]

    first_data_col = next((i for i, f in enumerate(header)
                           if re.fullmatch(r"\d+(?:\.\d+)?", f)), None)
    if first_data_col is None:
        raise ValueError(f"no numeric wavelength labels in the header of {path.name}")
    wavelengths = [int(float(f)) for f in header[first_data_col:]
                   if re.fullmatch(r"\d+(?:\.\d+)?", f)]

    scale = REFERENCE_PATH_MM / PATH_LENGTH_MM
    means, spreads = {}, {}
    for line in lines[header_idx + 1:]:
        parts = [p.strip() for p in line.split(delimiter)]
        if not re.fullmatch(r"[A-Z]+0?\d+", parts[0] or ""):
            continue
        values = []
        for token in parts[first_data_col:]:
            if delimiter == ";":
                token = token.replace(",", ".")
            try:
                values.append(float(token) * scale)
            except ValueError:
                # the reader writes "-" where a well returned no usable reading
                values.append(None)
        values = values[:len(wavelengths)]

        grouped = {}
        for wl, v in zip(wavelengths, values):
            if v is not None:
                grouped.setdefault(wl, []).append(v)
        well = normalise_well(parts[0])
        means[well] = {wl: statistics.fmean(v) for wl, v in grouped.items()}
        spreads[well] = {
            wl: (statistics.stdev(v) / statistics.fmean(v) * 100
                 if len(v) > 1 and statistics.fmean(v) else 0.0)
            for wl, v in grouped.items()
        }
    if not means:
        raise ValueError(f"no well rows found in {path.name}")
    return means, spreads


def read_plate_map(path):
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {name: header.index(name) for name in
           ("quencher_id", "name", "c_Stock / mM", "well_pos")}
    out = []
    for row in rows[1:]:
        if row[idx["quencher_id"]] is None:
            continue
        out.append({
            "quencher_id": int(row[idx["quencher_id"]]),
            "c_stock_mM": float(row[idx["c_Stock / mM"]]),
            "well": normalise_well(row[idx["well_pos"]]),
        })
    return out


def main():
    scan_path = RAW / SCAN
    map_path = RAW / PLATE_MAP
    for p in (scan_path, map_path):
        if not p.exists():
            sys.exit(f"missing input: {p}")

    means, spreads = read_scan(scan_path)
    quenchers = read_plate_map(map_path)

    missing_blanks = [w for w in BLANK_WELLS if w not in means]
    if missing_blanks:
        raise ValueError(f"blank wells absent from the scan: {missing_blanks}")
    blank = {wl: statistics.fmean([means[w][wl] for w in BLANK_WELLS]) for wl in WAVELENGTHS}

    columns = (["quencher_id", "well", "c_stock_mM"]
               + [f"epsilon_{wl}_mM-1cm-1" for wl in WAVELENGTHS]
               + ["max_read_rsd_percent"])

    rows = []
    for q in quenchers:
        well = q["well"]
        if well not in means:
            raise ValueError(f"quencher {q['quencher_id']} maps to well {well}, absent from the scan")
        row = {"quencher_id": q["quencher_id"],
               "well": well, "c_stock_mM": q["c_stock_mM"]}
        for wl in WAVELENGTHS:
            absorbance = means[well][wl] - blank[wl]
            row[f"epsilon_{wl}_mM-1cm-1"] = sig(absorbance / q["c_stock_mM"])
        row["max_read_rsd_percent"] = sig(max(spreads[well][wl] for wl in WAVELENGTHS))
        rows.append(row)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)

    worst = max(rows, key=lambda r: r["max_read_rsd_percent"])
    print(f"{len(rows)} quenchers x {len(WAVELENGTHS)} wavelengths -> {OUT.name}")
    print(f"  path length {PATH_LENGTH_MM} mm, scaled to {REFERENCE_PATH_MM} mm")
    print(f"  blanked against {len(BLANK_WELLS)} solvent wells ({BLANK_WELLS[0]}-{BLANK_WELLS[-1]})")
    print(f"  worst repeat-read spread: {worst['max_read_rsd_percent']}% "
          f"(Q{worst['quencher_id']})")


if __name__ == "__main__":
    main()
