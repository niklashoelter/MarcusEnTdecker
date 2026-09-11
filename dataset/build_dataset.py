# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""
Processed tables -> one row per photocatalyst-quencher pair.

Reads   ../data/experimental_data_processed/stern_volmer_fits.csv
        ../data/experimental_data_processed/photocatalyst_properties_exp.xlsx
        ../data/experimental_data_processed/quencher_properties_exp.xlsx
        ../data/computational_data_processed/photocatalyst_properties_calc.xlsx
        ../data/computational_data_processed/quencher_properties_calc.xlsx
Writes  pairs.csv

Nothing is read from outside the processed folders of ../data/, and no file is copied in
as an import side effect. The four workbooks are generated, not curated: each column's
`Column key` entry names the intermediate table in the raw folders it was copied from,
and from there every value can be followed back to the instrument file or job output it
came from, one step at a time.

THE CENTRAL QUANTITY

    k_q = K_SV / tau_0

K_SV per pair from the Stern-Volmer fits, tau_0 per photocatalyst from the TCSPC
traces. Both measured here.

WHY THE LABEL IS ON K_SV AND NOT k_q
    K_SV is what the experiment measures. k_q = K_SV / tau_0 is derived, and dividing by
    a photocatalyst-specific lifetime is what used to make the ground truth incoherent:
    at a fixed k_q cutoff of 1e7 the Stern-Volmer detection floor, 25/tau_0, sits ABOVE
    the cutoff for 19 of the 26 photocatalysts -- 33x above it for PC10. For those, "did
    not quench" did not mean k_q < 1e7; it meant nothing was shown above a floor well
    above 1e7, and 465 pairs could not be labelled at all.

    On K_SV the detection floor is the same number for every photocatalyst, so every
    pair is labelled by the same criterion and the indeterminate class does not arise.

    The cost is real and is not hidden: tau_0 spans 86x across the screen, so a fixed
    K_SV cutoff does correspond to different rate constants for different
    photocatalysts. That is the honest trade -- a threshold on what was measured, whose
    conversion to a rate varies, rather than a threshold on a derived rate that half the
    screen could not have detected. k_q, log10_k_q and k_q_floor are all still carried,
    because the rate is what the theories predict and what the rate-level analysis uses.

TWO LABELS, AND WHICH IS WHICH
    `quenches` is the label: K_SV > 10, or K_SV <= 10 with a slope significant across
    the points the fit used. `quenches_strict` is the published criterion, K_SV >= 25
    and r2 >= 0.90, carried so the stricter reading can be scored as a sensitivity
    check. Both come straight from stern_volmer_fits.csv, which explains the difference.

CENSORING IS CARRIED, NOT DROPPED
    A pair that does not clear the label is a one-sided observation, not missing data,
    and every one is in this table with `censored = 1` and its fitted K_SV intact,
    negative slopes included. `k_q_floor` gives the smallest rate constant that pair
    could have shown, `K_SV_FLOOR / tau_0`, which is the number a survival analysis
    needs and which varies by photocatalyst precisely because tau_0 does.

    Dropping them would bias every fit in the paper towards the pairs that happened to
    quench.

BOUNDS ARE CARRIED TOO
    Nineteen quenchers have a redox potential recorded as a limit rather than a value.
    `E_Ox_bound` / `E_Red_bound` say so. Anything that rules out the single-electron
    transfer channel by comparing potentials has to look at those columns, because
    treating `> 2.0` as `2.0` makes oxidation look reachable when the measurement only
    established that it is not reachable below 2.0 V.

THE DRIVING FORCES
    Both channels a triplet photocatalyst can quench through, per pair, all
    quencher-centric and all negative when exergonic:

        dG_EnT     = E_T(Q) - E_T(PC)
        dG_SET_ox  = F * (E_Ox(Q) - E(PC*/PC-))     Q oxidized, PC* is the oxidant
        dG_SET_red = F * (E(PC+/PC*) - E_Red(Q))    Q reduced,  PC* is the reductant

    with F = 23.0605 kcal mol-1 V-1. `dG_SET` is whichever of the two is more
    accessible, and `dG_SET_channel` says which, because "the SET driving force" is
    meaningless without knowing which way the electron went.

    Potentials are named by their couple rather than by direction. E(PC*/PC-) is the
    potential at which PC* *oxidizes* a substrate -- the reduction potential of the
    excited state. Naming it "the excited-state oxidation potential" is how the sign
    gets flipped.

    Both SET columns are computed twice, from the measured quencher potentials
    (`_exp`) and from the computed ones (`_calc`), and the two are never merged. They
    are different quantities with different provenance, and which one an analysis used
    is part of what it has to report.

BOUNDED POTENTIALS PROPAGATE INTO BOUNDED DRIVING FORCES
    Nineteen quenchers have a redox potential recorded as a limit. Both limits point
    the same way once they reach a driving force: an oxidation recorded as `> 2.0` puts
    a floor under dG_SET_ox, and a reduction recorded as `< -2.10` puts a floor under
    dG_SET_red. So `dG_SET_ox_bound = lower` means SET is *at least* this endergonic --
    the true value can only be worse. That is a usable statement, and it is the only
    honest one; reading the number at face value would understate how inaccessible the
    channel is, in the direction that wrongly excludes pairs from the energy-transfer
    set.

WHAT IS STILL NOT HERE
    The photocatalyst reorganization energies are carried but no combined lambda is
    formed. Marcus theory needs the pair's lambda, and how the two partners' terms
    combine is a modelling decision that belongs in benchmark/, not in the table.
"""

import collections
import csv
import math
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
EXPERIMENTAL = ROOT / "data" / "experimental_data_processed"
COMPUTATIONAL = ROOT / "data" / "computational_data_processed"
OUT = Path(__file__).resolve().parent / "pairs.csv"

# Row 1 of every property workbook names the column blocks; the headers are on row 2.
# See data/properties_workbook.py, whose read_workbook() this mirrors.
WORKBOOK_HEADER_ROW = 2

# The rate constant the old label used. Kept only to report, in the summary below, how
# far the K_SV label and that k_q cutoff disagree; nothing is classified on it.
K_Q_THRESHOLD = 1.0e7

# The K_SV a pair has to clear to count as detected, and so the left-censoring limit for
# anything fitted on K_SV. Set where the label is set, in stern_volmer_fits.py, and equal
# to K_SV_LABEL there. Fixed in K_SV; photocatalyst-specific once divided by tau_0, which
# is why k_q_floor has to be carried per row.
K_SV_FLOOR = 10.0

SIGFIGS = 6

# Faraday constant in the units the driving forces are reported in.
F_KCAL_PER_V = 23.060548

# Photocatalysts measured but excluded from the dataset, with the reason.
#
# This is a scientific decision taken outside this script, recorded here so that the
# 1296 rows the Stern-Volmer fits would otherwise produce and the 1248 written below
# differ by something a reader can see and disagree with. It is NOT a data-quality
# filter and no threshold produces it: PC_24 is unremarkable on both fit quality and
# lifetime provenance, and three photocatalysts that are kept have worse median r2.
EXCLUDED_PHOTOCATALYSTS = {
    24: "discarded from the screen; marked 'discard' in the photocatalyst record",
}

# The quencher workbook writes a bounded potential the way the lab records do, as the
# inequality it is read with; this table names the side of the limit instead.
BOUND_NAME = {">": "lower", "<": "upper"}

# Wavelengths at which an extinction coefficient was measured, for matching a quencher's
# absorbance to the wavelength its photocatalyst was excited at.
EPSILON_NM = [380, 390, 400, 405, 415, 425, 450, 475, 490, 495, 500,
              510, 520, 580, 590, 600, 625, 645]

COLUMNS = [
    # identity
    "pc_id", "pc_name", "pc_compound", "quencher_id", "quencher_name", "quencher_compound",
    "quencher_smiles",
    # the measurement
    "K_SV / M-1", "r2", "n_points_fitted", "n_points_measured", "fit_method", "fit_note",
    "tau0 / ns", "tau0_rsd", "tau0_consistent",
    # the rate
    "k_q / M-1 s-1", "log10_k_q", "quenches", "quenches_strict", "censored",
    "k_q_floor / M-1 s-1", "K_SV_floor / M-1",
    # quencher, experimental
    "q_E_Ox_vs_SCE_exp / V", "q_E_Ox_bound", "q_E_Red_vs_SCE_exp / V", "q_E_Red_bound",
    "q_redox_source_type", "q_redox_doi",
    "q_epsilon_at_pc_excitation / mM-1 cm-1", "pc_excitation / nm",
    # quencher, computed
    "q_E_T_dft / kcal mol-1", "q_lambda_S_T_dft / kcal mol-1", "q_lambda_T_S_dft / kcal mol-1",
    "q_E_Ox_vs_SCE_calc / V", "q_E_Red_vs_SCE_calc / V",
    "q_dvte_mu_md / kcal mol-1", "q_dvte_std_md / kcal mol-1",
    # photocatalyst, computed
    "pc_E_T_dft / kcal mol-1",
    "pc_lambda_S0_to_T1_dft / kcal mol-1", "pc_lambda_T1_to_S0_dft / kcal mol-1",
    "pc_E_PCstar_PCminus / V", "pc_E_PCplus_PCstar / V",
    # driving forces
    "dG_EnT / kcal mol-1",
    "dG_SET_ox_exp / kcal mol-1", "dG_SET_ox_exp_bound",
    "dG_SET_red_exp / kcal mol-1", "dG_SET_red_exp_bound",
    "dG_SET_exp / kcal mol-1", "dG_SET_exp_bound", "dG_SET_exp_channel",
    "dG_SET_ox_calc / kcal mol-1", "dG_SET_red_calc / kcal mol-1",
    "dG_SET_calc / kcal mol-1", "dG_SET_calc_channel",
]


def sig(x, n=SIGFIGS):
    if x in (None, ""):
        return ""
    return 0.0 if float(x) == 0 else float(f"{float(x):.{n}g}")


def load(path, key):
    """Read a property workbook into {key: row}. Missing file is fatal, not skipped.

    Blank cells come back as "", so a row reads exactly as the CSV it replaced did, and the
    empty columns that separate the blocks are dropped.
    """
    if not path.exists():
        sys.exit(f"missing input: {path}\nrun the data stage first: uv run reproduce.py data")
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(min_row=WORKBOOK_HEADER_ROW, values_only=True)
    header = next(rows)
    table = {}
    for values in rows:
        row = {h: ("" if v is None else v) for h, v in zip(header, values) if h}
        if row.get(key) not in ("", None):
            table[int(row[key])] = row
    return table


def number(row, field):
    """A numeric field, or None when absent or blank."""
    if row is None:
        return None
    value = row.get(field, "")
    return float(value) if value not in ("", None) else None


# label_class() used to live here. It sorted pairs into positive / negative /
# indeterminate because the k_q cutoff and the per-photocatalyst detection floor
# disagreed, leaving 465 pairs the experiment could not speak to. Labelling on K_SV
# removes the disagreement -- one floor, the same for every photocatalyst -- so every
# pair is now either detected or censored and the third class has nothing in it.


def better_set_channel(ox, red):
    """The more accessible of the two SET channels: (value, 'ox'|'red'), or (None, '').

    "More accessible" is the less positive driving force. If only one channel has the
    potentials it needs, that one is returned -- a missing channel is not evidence that
    the other one is the preferred route, so the column says which one it is reporting.
    """
    available = [(value, name) for value, name in ((ox, "ox"), (red, "red"))
                 if value is not None]
    if not available:
        return None, ""
    return min(available)


def nearest_epsilon(absorption_row, excitation_nm):
    """Extinction coefficient at the wavelength closest to the excitation wavelength.

    Excitation was at 380, 405, 425 or 450 nm and every one of those was measured
    directly, so `nearest` is an exact match in practice. It is written as a search
    rather than a lookup so that a photocatalyst excited somewhere else still resolves,
    and `epsilon_offset_nm` would be the thing to add if that ever happens.
    """
    if absorption_row is None or excitation_nm is None:
        return None
    nm = min(EPSILON_NM, key=lambda w: abs(w - excitation_nm))
    return number(absorption_row, f"epsilon_{nm} / mM-1 cm-1")


def main():
    fits_path = EXPERIMENTAL / "stern_volmer_fits.csv"
    if not fits_path.exists():
        sys.exit(f"missing input: {fits_path}\nrun the data stage first")

    # One workbook per compound class and side. The lifetimes, redox records and
    # extinction coefficients share a row, as do the adiabatic DFT and the DvTE columns,
    # so each name below is the same table seen for one purpose.
    photocatalysts = tau0 = load(EXPERIMENTAL / "photocatalyst_properties_exp.xlsx", "pc_id")
    quenchers = absorption = load(EXPERIMENTAL / "quencher_properties_exp.xlsx", "quencher_id")
    dft = dvte = load(COMPUTATIONAL / "quencher_properties_calc.xlsx", "quencher_id")
    pc_dft = load(COMPUTATIONAL / "photocatalyst_properties_calc.xlsx", "pc_id")

    rows, no_tau0, quench_count, censored_count = [], set(), 0, 0
    strict_count = disagrees_with_kq = 0
    excluded_rows = 0

    with fits_path.open(encoding="utf-8", newline="") as fh:
        for fit in csv.DictReader(fh):
            pc_id, q_id = int(fit["pc_id"]), int(fit["quencher_id"])
            if pc_id in EXCLUDED_PHOTOCATALYSTS:
                excluded_rows += 1
                continue

            lifetime = number(tau0.get(pc_id), "tau0 / ns")
            if lifetime is None:
                no_tau0.add(pc_id)

            k_sv = number(fit, "K_SV")
            # tau_0 is in ns; k_q comes out in M-1 s-1.
            k_q = (k_sv / (lifetime * 1e-9)) if (k_sv is not None and lifetime) else None
            floor = (K_SV_FLOOR / (lifetime * 1e-9)) if lifetime else None

            # Both labels are set in stern_volmer_fits.py, on the fit and its points.
            # Nothing about them is re-decided here.
            quenches = fit["quenches"] == "1"
            quenches_strict = fit["quenching_observed"] == "True"
            # Failing the label makes the pair left-censored: K_SV is somewhere below
            # the floor, not zero and not missing.
            censored = not quenches
            quench_count += quenches
            strict_count += quenches_strict
            censored_count += censored
            if quenches != (k_q is not None and k_q > K_Q_THRESHOLD):
                disagrees_with_kq += 1

            pc = photocatalysts.get(pc_id)
            q = quenchers.get(q_id)
            excitation = number(pc, "excitation / nm")

            pcc = pc_dft.get(pc_id)
            pc_e_t = number(pcc, "E_T / kcal mol-1")
            pc_star_ox = number(pcc, "E_PCstar_PCminus / V")   # PC* as the oxidant
            pc_star_red = number(pcc, "E_PCplus_PCstar / V")   # PC* as the reductant

            q_e_t = number(dft.get(q_id), "E_T / kcal mol-1")
            dG_ent = (q_e_t - pc_e_t) if None not in (q_e_t, pc_e_t) else None

            def set_driving_forces(e_ox, e_red):
                """(dG_SET_ox, dG_SET_red) from one pair of quencher potentials."""
                ox = (F_KCAL_PER_V * (e_ox - pc_star_ox)
                      if None not in (e_ox, pc_star_ox) else None)
                red = (F_KCAL_PER_V * (pc_star_red - e_red)
                       if None not in (e_red, pc_star_red) else None)
                return ox, red

            dG_ox_exp, dG_red_exp = set_driving_forces(
                number(q, "E_Ox_vs_SCE_exp / V"), number(q, "E_Red_vs_SCE_exp / V"))
            dG_ox_calc, dG_red_calc = set_driving_forces(
                number(dft.get(q_id), "E_Ox_vs_SCE_calc / V"),
                number(dft.get(q_id), "E_Red_vs_SCE_calc / V"))

            # A bound on either potential becomes a lower bound on the driving force:
            # the quencher was not oxidized below the recorded value, so SET is at
            # least this endergonic. Both directions land on `lower` for that reason.
            ox_bound = "lower" if (q or {}).get("E_Ox_bound") else ""
            red_bound = "lower" if (q or {}).get("E_Red_bound") else ""

            dG_set_exp, exp_channel = better_set_channel(dG_ox_exp, dG_red_exp)
            dG_set_calc, calc_channel = better_set_channel(dG_ox_calc, dG_red_calc)
            set_bound = {"ox": ox_bound, "red": red_bound}.get(exp_channel, "")

            rows.append({
                "pc_id": pc_id,
                "pc_name": (pc or {}).get("pc_name", ""),
                "pc_compound": (pc or {}).get("compound", ""),
                "quencher_id": q_id,
                "quencher_name": (q or {}).get("quencher_name", ""),
                "quencher_compound": (q or {}).get("compound", ""),
                "quencher_smiles": (q or {}).get("smiles", ""),

                "K_SV / M-1": sig(k_sv),
                "r2": sig(number(fit, "r2")),
                "n_points_fitted": fit["num_points"],
                "n_points_measured": fit["n_measured"],
                "fit_method": fit["fit_method"],
                "fit_note": fit["note"],
                "tau0 / ns": sig(lifetime),
                "tau0_rsd": sig(number(tau0.get(pc_id), "tau0_rsd")),
                "tau0_consistent": (tau0.get(pc_id) or {}).get("tau0_consistent", ""),

                "k_q / M-1 s-1": sig(k_q),
                "log10_k_q": sig(math.log10(k_q)) if (k_q and k_q > 0) else "",
                "quenches": int(quenches),
                "quenches_strict": int(quenches_strict),
                "censored": int(censored),
                "k_q_floor / M-1 s-1": sig(floor),
                "K_SV_floor / M-1": sig(K_SV_FLOOR),

                "q_E_Ox_vs_SCE_exp / V": sig(number(q, "E_Ox_vs_SCE_exp / V")),
                "q_E_Ox_bound": BOUND_NAME.get((q or {}).get("E_Ox_bound", ""), ""),
                "q_E_Red_vs_SCE_exp / V": sig(number(q, "E_Red_vs_SCE_exp / V")),
                "q_E_Red_bound": BOUND_NAME.get((q or {}).get("E_Red_bound", ""), ""),
                "q_redox_source_type": (q or {}).get("redox_source_type", ""),
                "q_redox_doi": (q or {}).get("redox_doi", ""),
                "q_epsilon_at_pc_excitation / mM-1 cm-1":
                    sig(nearest_epsilon(absorption.get(q_id), excitation)),
                "pc_excitation / nm": sig(excitation),

                "q_E_T_dft / kcal mol-1": sig(q_e_t),
                "q_lambda_S_T_dft / kcal mol-1":
                    sig(number(dft.get(q_id), "lambda_S_T / kcal mol-1")),
                "q_lambda_T_S_dft / kcal mol-1":
                    sig(number(dft.get(q_id), "lambda_T_S / kcal mol-1")),
                "q_E_Ox_vs_SCE_calc / V": sig(number(dft.get(q_id), "E_Ox_vs_SCE_calc / V")),
                "q_E_Red_vs_SCE_calc / V": sig(number(dft.get(q_id), "E_Red_vs_SCE_calc / V")),
                "q_dvte_mu_md / kcal mol-1": sig(number(dvte.get(q_id), "dvte_mu / kcal mol-1")),
                "q_dvte_std_md / kcal mol-1": sig(number(dvte.get(q_id), "dvte_std / kcal mol-1")),

                "pc_E_T_dft / kcal mol-1": sig(pc_e_t),
                "pc_lambda_S0_to_T1_dft / kcal mol-1":
                    sig(number(pcc, "lambda_S0_to_T1 / kcal mol-1")),
                "pc_lambda_T1_to_S0_dft / kcal mol-1":
                    sig(number(pcc, "lambda_T1_to_S0 / kcal mol-1")),
                "pc_E_PCstar_PCminus / V": sig(pc_star_ox),
                "pc_E_PCplus_PCstar / V": sig(pc_star_red),

                "dG_EnT / kcal mol-1": sig(dG_ent),
                "dG_SET_ox_exp / kcal mol-1": sig(dG_ox_exp),
                "dG_SET_ox_exp_bound": ox_bound if dG_ox_exp is not None else "",
                "dG_SET_red_exp / kcal mol-1": sig(dG_red_exp),
                "dG_SET_red_exp_bound": red_bound if dG_red_exp is not None else "",
                "dG_SET_exp / kcal mol-1": sig(dG_set_exp),
                "dG_SET_exp_bound": set_bound,
                "dG_SET_exp_channel": exp_channel,
                "dG_SET_ox_calc / kcal mol-1": sig(dG_ox_calc),
                "dG_SET_red_calc / kcal mol-1": sig(dG_red_calc),
                "dG_SET_calc / kcal mol-1": sig(dG_set_calc),
                "dG_SET_calc_channel": calc_channel,
            })

    rows.sort(key=lambda row: (row["pc_id"], row["quencher_id"]))
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    pcs = {row["pc_id"] for row in rows}
    qs = {row["quencher_id"] for row in rows}
    print(f"{len(rows)} pairs -> {OUT.name}   ({len(pcs)} photocatalysts x {len(qs)} quenchers)")
    for pc_id, reason in EXCLUDED_PHOTOCATALYSTS.items():
        print(f"  excluded PC{pc_id}: {reason}")
    if excluded_rows:
        print(f"  {excluded_rows} fitted pairs dropped by that exclusion, not by any threshold")
    if len(rows) != len(pcs) * len(qs):
        print(f"  WARNING: not a complete grid, expected {len(pcs) * len(qs)}")
    if no_tau0:
        print(f"  WARNING: no lifetime for photocatalysts {sorted(no_tau0)} — k_q is empty there")

    rates = [row["k_q / M-1 s-1"] for row in rows if row["k_q / M-1 s-1"] != ""]
    print(f"  quenches (K_SV label):       {quench_count} yes, {len(rows) - quench_count} no")
    print(f"  quenches_strict (published): {strict_count} yes, {len(rows) - strict_count} no")
    print(f"  left-censored (below the K_SV floor of {K_SV_FLOOR:g}): {censored_count}")
    print(f"  k_q computed for {len(rates)} pairs; the old k_q > {K_Q_THRESHOLD:.0e} cutoff "
          f"disagrees with the label on {disagrees_with_kq}")
    if rates:
        floors = [row["k_q_floor / M-1 s-1"] for row in rows if row["k_q_floor / M-1 s-1"] != ""]
        print(f"  detection floor ranges {min(floors):.2e} to {max(floors):.2e} M-1 s-1 "
              f"({max(floors) / min(floors):.0f}x, because tau_0 does) -- which is why the "
              f"label is not set on k_q")

    sources = collections.Counter(row["q_redox_source_type"] for row in rows)
    print(f"  redox provenance: " + ", ".join(f"{v} {k}" for k, v in sources.most_common() if k))

    for label, field in (("quencher SMILES", "quencher_smiles"),
                         ("quencher DFT E_T", "q_E_T_dft / kcal mol-1"),
                         ("quencher DvTE", "q_dvte_mu_md / kcal mol-1"),
                         ("experimental E_Ox", "q_E_Ox_vs_SCE_exp / V"),
                         ("photocatalyst DFT E_T", "pc_E_T_dft / kcal mol-1"),
                         ("dG_EnT", "dG_EnT / kcal mol-1"),
                         ("dG_SET from measured potentials", "dG_SET_exp / kcal mol-1"),
                         ("dG_SET from computed potentials", "dG_SET_calc / kcal mol-1")):
        have = sum(1 for row in rows if row[field] != "")
        print(f"  {label}: {have}/{len(rows)}")

    exergonic = sum(1 for row in rows
                    if row["dG_EnT / kcal mol-1"] != "" and row["dG_EnT / kcal mol-1"] < 0)
    bounded = sum(1 for row in rows if row["dG_SET_exp_bound"])
    channels = collections.Counter(row["dG_SET_exp_channel"] for row in rows)
    print(f"\n  dG_EnT exergonic for {exergonic} pairs")
    print(f"  dG_SET (measured) preferred channel: "
          + ", ".join(f"{v} {k}" for k, v in channels.most_common() if k))
    print(f"  ...of which {bounded} rest on a bounded potential and are lower bounds only")


if __name__ == "__main__":
    main()
