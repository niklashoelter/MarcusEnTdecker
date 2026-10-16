"""
Python API for triplet energy transfer rates from structure.

    from ml.predict import RatePredictor, photocatalyst_properties

    predictor = RatePredictor()
    result = predictor.predict(photocatalyst=photocatalyst_properties(pc_table, 1),
                               quencher=Q_SMILES)

    result.dG_kcal          driving force
    result.dG_barrier_kcal  Marcus activation barrier
    result.k_ent            rate constant, M-1 s-1
    result.t_pc_kcal        lifetime barrier t_PC, from the photocatalyst's tau0
    result.two_gate         dG <= 0 and dG_barrier <= t_PC; None without tau0
    result.uncertainty      from the mixture density network

`ml.cli` wraps these calls as the `ent-ml` command.

    ml.prediction.predict(smiles)   graph -> E_T, lambda_S_T, lambda_T_S
    ml.energy.marcus_ent            properties -> dG, dG barrier
    ml.energy.two_gate              tau0 -> t_PC, two-gate verdict, k_-d/nu
    RatePredictor._encounter_rate   barrier -> k

Rate:
    k_EnT = k_diff / (1 + exp(dG/RT) + (k_-d/nu) * exp(dG_barrier/RT))
    k_-d = k_diff / K_A, K_A the Eigen encounter constant at contact distance 7 A,
    nu = k_B T / h. The Eyring rate of the barrier is returned as `k_eyring`.

Two-gate: True when dG <= 0 and dG_barrier <= t_PC, the barrier at which k * tau0 reaches
K_SV,crit = 10 M-1. Needs `tau0_ns` (`tau0 / ns` column of the table); otherwise
`t_pc_kcal` and `two_gate` are None.

Element vocabulary: H, B, C, N, O, F, Si, P, S, Cl, Se, Br, Sn, I; Ir and Ru raise
`UnknownAtomError`. A metal-complex photocatalyst is passed as DFT properties,
`photocatalyst_properties(source, pc_id)`. Each partner is a SMILES (model) or
`MolecularProperties` (used as given). `uncertainty` covers only the model-predicted
partners, named in `uncertainty_covers`.
"""

from __future__ import annotations

import csv
import math
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence, Union

# Property workbooks: headers on row 2.
WORKBOOK_HEADER_ROW = 2

#: Required photocatalyst columns besides `new_id`.
PC_COLUMNS = ("E_T / kcal mol-1", "lambda_S0_T1 / kcal mol-1",
              "lambda_T1_S0 / kcal mol-1")

#: Optional photocatalyst column: excited-state lifetime, for the two-gate verdict.
PC_TAU0 = "tau0 / ns"

# Gas constant in kcal mol-1 K-1.
R_KCAL = 1.987204259e-3

# Acetonitrile diffusion limit, M-1 s-1.
K_DIFF_ACETONITRILE = 2.0e10

DEFAULT_TEMPERATURE_K = 298.15


@dataclass(frozen=True)
class MolecularProperties:
    """What the model predicts for one structure. Energies in kcal/mol."""

    smiles: str
    E_T: float
    lambda_S_T: float
    lambda_T_S: float
    # Uncertainties: None for DFT properties.
    E_T_uncertainty: float | None = None
    lambda_S_T_uncertainty: float | None = None
    lambda_T_S_uncertainty: float | None = None
    spin_density: list[float] = field(default_factory=list)
    source: str = "model"
    # Measured excited-state lifetime of a photocatalyst, ns; None when unknown.
    tau0_ns: float | None = None


#: Either a structure the model can encode, or properties already in hand.
Partner = Union[str, "MolecularProperties"]


def _rows(source: Any) -> list[dict]:
    """A photocatalyst table as a list of row dicts."""
    if source is None:
        raise ValueError(
            "photocatalyst_properties needs a photocatalyst table: a path to an .xlsx "
            "or .csv, a pandas DataFrame, or a mapping. One row per photocatalyst, "
            "with a `new_id` column and " + ", ".join(f"`{c}`" for c in PC_COLUMNS) + "."
        )

    if hasattr(source, "to_dict") and hasattr(source, "columns"):  # a DataFrame
        return source.to_dict("records")

    if isinstance(source, Mapping):
        # Either a whole table keyed by id, or one row already picked out of one.
        values = list(source.values())
        if values and all(isinstance(v, Mapping) for v in values):
            return [{"new_id": key, **row} for key, row in source.items()]
        return [dict(source)]

    path = Path(source)
    if not path.exists():
        raise FileNotFoundError(f"no photocatalyst table at {path}")

    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle))

    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    values = sheet.iter_rows(min_row=WORKBOOK_HEADER_ROW, values_only=True)
    header = next(values)
    return [dict(zip(header, v)) for v in values]


def photocatalyst_properties(source: Any = None,
                             pc_id: int | str | None = None) -> MolecularProperties:
    """DFT properties of one photocatalyst, as `MolecularProperties`.

    `source`: path to an .xlsx (headers on row 2) or .csv, a DataFrame, a mapping id -> row,
    or one row. `pc_id`: the row's `new_id`; optional for a one-row table.
    `lambda_S0_T1` (E_S0 at the T1 geometry - E_S0) -> `lambda_S_T`,
    `lambda_T1_S0` (E_T1 at the S0 geometry - E_T1) -> `lambda_T_S`,
    `tau0 / ns` (optional) -> `tau0_ns`.
    """
    rows = _rows(source)

    if pc_id is None:
        if len(rows) != 1:
            raise ValueError(
                f"the photocatalyst table has {len(rows)} rows; say which one with pc_id"
            )
        row = rows[0]
    else:
        # An id read from a .csv or a command line is a string, one from an .xlsx an int.
        row = next((r for r in rows if str(r.get("new_id")) == str(pc_id)), None)
        if row is None:
            raise KeyError(f"no photocatalyst PC{pc_id} in the table")

    identifier = row.get("new_id", pc_id)
    name = f"PC{identifier}" if identifier is not None else "photocatalyst"

    absent = [c for c in PC_COLUMNS if row.get(c) in (None, "")]
    if absent:
        raise ValueError(
            f"{name} has no {', '.join(absent)} -- its DFT jobs are incomplete. "
            f"It cannot be used for a rate prediction."
        )

    tau0 = row.get(PC_TAU0)
    if tau0 in (None, "") or (isinstance(tau0, float) and math.isnan(tau0)):
        tau0 = None  # a blank cell from a .csv, an .xlsx or a DataFrame
    return MolecularProperties(
        smiles=name,
        E_T=float(row[PC_COLUMNS[0]]),
        lambda_S_T=float(row[PC_COLUMNS[1]]),
        lambda_T_S=float(row[PC_COLUMNS[2]]),
        source="dft",
        tau0_ns=None if tau0 is None else float(tau0),
    )


@dataclass(frozen=True)
class RatePrediction:
    """A predicted rate constant for one photocatalyst-quencher pair."""

    photocatalyst: str
    quencher: str

    dG_kcal: float
    dG_barrier_kcal: float
    dG_barrier_reverse_kcal: float

    k_ent: float
    k_eyring: float

    t_pc_kcal: float | None
    two_gate: bool | None

    uncertainty: float | None
    uncertainty_covers: str
    symmetric_approximation: bool

    photocatalyst_properties: MolecularProperties
    quencher_properties: MolecularProperties

    @property
    def log10_k_ent(self) -> float:
        return math.log10(self.k_ent) if self.k_ent > 0 else float("-inf")


class RatePredictor:
    """EnT rate constants for photocatalyst-quencher pairs.

    The model loads on the first prediction; properties are cached per SMILES.
    `ensemble=True`: properties from `ml.ensemble`.
    """

    def __init__(
        self,
        temperature: float = DEFAULT_TEMPERATURE_K,
        k_diff: float = K_DIFF_ACETONITRILE,
        ensemble: bool = False,
    ):
        self.temperature = temperature
        self.k_diff = k_diff
        self.ensemble = ensemble
        self._cache: dict[str, MolecularProperties] = {}

    # -- properties -------------------------------------------------------------

    def properties(self, smiles: Sequence[str]) -> list[MolecularProperties]:
        """E_T and both reorganization energies per SMILES; uncached ones in one batch."""
        wanted = list(smiles)
        missing = [s for s in dict.fromkeys(wanted) if s not in self._cache]

        if missing:
            if self.ensemble:
                from ml.ensemble import predict as _predict
            else:
                from ml.prediction import predict as _predict

            frame = _predict(missing)
            for _, row in frame.iterrows():
                self._cache[row["smiles"]] = MolecularProperties(
                    smiles=row["smiles"],
                    E_T=float(row["triplet_energy_adiabatic"]),
                    E_T_uncertainty=float(row["triplet_energy_adiabatic_uncertainty"]),
                    lambda_S_T=float(row["lamda_S_T"]),
                    lambda_S_T_uncertainty=float(row["lamda_S_T_uncertainty"]),
                    lambda_T_S=float(row["lamda_T_S"]),
                    lambda_T_S_uncertainty=float(row["lamda_T_S_uncertainty"]),
                    spin_density=list(row["spin_map"]),
                    source="ensemble" if self.ensemble else "model",
                )

        return [self._cache[s] for s in wanted]

    # -- rates ------------------------------------------------------------------

    def predict(self, photocatalyst: Partner, quencher: Partner) -> RatePrediction:
        """Rate constant for one pair; each partner a SMILES or `MolecularProperties`."""
        return self.predict_many([(photocatalyst, quencher)])[0]

    def predict_many(self, pairs: Iterable[tuple[Partner, Partner]]) -> list[RatePrediction]:
        """Rate constants for many pairs; each distinct SMILES is predicted once."""
        pairs = list(pairs)
        if not pairs:
            return []

        smiles = [p for pair in pairs for p in pair if isinstance(p, str)]
        self.properties(list(dict.fromkeys(smiles)))

        def resolve(partner: Partner) -> MolecularProperties:
            return self._cache[partner] if isinstance(partner, str) else partner

        return [self._rate(resolve(pc), resolve(q)) for pc, q in pairs]

    def _rate(self, pc: MolecularProperties, q: MolecularProperties) -> RatePrediction:
        from ml.energy.eyring import eyring_equation
        from ml.energy.marcus_ent import (
            calculate_driving_force_and_barrier,
        )
        from ml.energy.two_gate import lifetime_barrier, quenches as two_gate_quenches

        dG, barrier, barrier_reverse = calculate_driving_force_and_barrier(
            pc_adiabatic_E_T=pc.E_T,
            pc_lambda_S_T=pc.lambda_S_T,
            pc_lambda_T_S=pc.lambda_T_S,
            quencher_E_T=q.E_T,
            quencher_lambda_S_T=q.lambda_S_T,
            quencher_lambda_T_S=q.lambda_T_S,
        )

        k_ent = self._encounter_rate(dG, barrier)
        t_pc = (None if pc.tau0_ns is None
                else lifetime_barrier(pc.tau0_ns, temperature=self.temperature))

        # E_T uncertainties of the model-predicted partners, summed in quadrature.
        estimated = [(name, p.E_T_uncertainty) for name, p in
                     (("photocatalyst", pc), ("quencher", q))
                     if p.E_T_uncertainty is not None]
        uncertainty = (math.sqrt(sum(u * u for _, u in estimated)) if estimated else None)
        covers = " and ".join(name for name, _ in estimated) or "neither partner"

        return RatePrediction(
            photocatalyst=pc.smiles,
            quencher=q.smiles,
            dG_kcal=float(dG),
            dG_barrier_kcal=float(barrier),
            dG_barrier_reverse_kcal=float(barrier_reverse),
            k_ent=k_ent,
            k_eyring=float(eyring_equation(barrier, T=self.temperature)),
            t_pc_kcal=t_pc,
            two_gate=None if t_pc is None else two_gate_quenches(dG, barrier, t_pc),
            uncertainty=uncertainty,
            uncertainty_covers=covers,
            # symmetric fallback of calculate_driving_force_and_barrier (no intersection)
            symmetric_approximation=self._is_symmetric(pc, q),
            photocatalyst_properties=pc,
            quencher_properties=q,
        )

    def _encounter_rate(self, dG: float, barrier: float) -> float:
        """k_diff / (1 + exp(dG/RT) + (k_-d/nu) exp(barrier/RT)), guarded against overflow."""
        from ml.energy.two_gate import escape_over_frequency

        rt = R_KCAL * self.temperature
        escape_over_nu = escape_over_frequency(temperature=self.temperature,
                                               k_diff=self.k_diff)

        def bounded_exp(x):
            # exponent capped at 700
            return math.exp(min(x / rt, 700.0))

        denominator = 1.0 + bounded_exp(dG) + escape_over_nu * bounded_exp(barrier)
        return self.k_diff / denominator

    @staticmethod
    def _is_symmetric(pc: MolecularProperties, q: MolecularProperties) -> bool:
        """Whether ml.energy.marcus_ent uses the symmetric approximation."""
        forward = pc.lambda_T_S + q.lambda_S_T
        reverse = pc.lambda_S_T + q.lambda_T_S
        dG = q.E_T - pc.E_T
        root = forward * reverse + (forward - reverse) * dG
        return root < 0 or abs(forward - reverse) < 1e-10
